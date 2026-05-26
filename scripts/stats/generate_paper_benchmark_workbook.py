#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
合并两份统计 Excel 为论文用单册 `paper_benchmark_tables.xlsx`：

- `comprehensive_benchmark_stats.xlsx`
- `benchmark_source_model_summary.xlsx`

各 sheet 内已按「节标题 + 多张表纵向堆叠」排版，可直接复制到论文。
"""
from __future__ import annotations

import sys
from copy import copy
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[2]
_SCRIPTS = _ROOT / "scripts"
for p in (str(_ROOT), str(_SCRIPTS)):
    if p not in sys.path:
        sys.path.insert(0, p)

from openpyxl import Workbook, load_workbook

REPORTS = _ROOT / "output/reports"
COMPREHENSIVE = REPORTS / "comprehensive_benchmark_stats.xlsx"
BENCHMARK = REPORTS / "benchmark_source_model_summary.xlsx"
OUTPUT = REPORTS / "paper_benchmark_tables.xlsx"

SHEET_ORDER = [
    "00_说明与覆盖",
    "表1_数据集概览",
    "表2_难度与区分度",
    "表3_模型得分",
    "表4_L1任务类型",
    "表5_模型排名与满分率",
    "表6_裁判与区分度",
    "附录_题级明细",
    "Charts",
]


def _copy_worksheet(source, target_wb, new_title: str) -> None:
    if new_title in target_wb.sheetnames:
        del target_wb[new_title]
    target = target_wb.create_sheet(new_title)

    for row in source.iter_rows():
        for cell in row:
            tc = target.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                tc.font = copy(cell.font)
                tc.border = copy(cell.border)
                tc.fill = copy(cell.fill)
                tc.number_format = cell.number_format
                tc.protection = copy(cell.protection)
                tc.alignment = copy(cell.alignment)

    for merged in source.merged_cells.ranges:
        target.merge_cells(str(merged))

    for col, dim in source.column_dimensions.items():
        target.column_dimensions[col].width = dim.width
    for row, dim in source.row_dimensions.items():
        target.row_dimensions[row].height = dim.height

    for img in getattr(source, "_images", []):
        target.add_image(copy(img), img.anchor)


def _append_meta_note(wb) -> None:
    """在说明 sheet 末尾追加模型汇总来源提示。"""
    if "00_说明与覆盖" not in wb.sheetnames:
        return
    ws = wb["00_说明与覆盖"]
    start = ws.max_row + 3
    ws.cell(start, 1, "【说明】模型排名 / 裁判一致性 / 满分率矩阵").font = copy(
        ws.cell(1, 1).font
    )
    ws.cell(start + 1, 1, "见本册「表5_模型排名与满分率」「表6_裁判与区分度」")


def run(
    comprehensive: Path = COMPREHENSIVE,
    benchmark: Path = BENCHMARK,
    output: Path = OUTPUT,
) -> None:
    if not comprehensive.is_file():
        raise FileNotFoundError(f"缺少: {comprehensive}")
    if not benchmark.is_file():
        raise FileNotFoundError(f"缺少: {benchmark}")

    dst = Workbook()
    dst.remove(dst.active)

    src_comp = load_workbook(comprehensive)
    for name in SHEET_ORDER:
        if name in src_comp.sheetnames and name not in ("表5_模型排名与满分率", "表6_裁判与区分度"):
            _copy_worksheet(src_comp[name], dst, name)

    src_bench = load_workbook(benchmark)
    for name in ("表5_模型排名与满分率", "表6_裁判与区分度"):
        if name in src_bench.sheetnames:
            _copy_worksheet(src_bench[name], dst, name)

    for i, name in enumerate(SHEET_ORDER):
        if name in dst.sheetnames:
            dst.move_sheet(name, offset=i - dst.sheetnames.index(name))

    _append_meta_note(dst)

    output.parent.mkdir(parents=True, exist_ok=True)
    dst.save(output)
    print(f"✓ 论文用总册: {output}")
    print(f"  sheets ({len(dst.sheetnames)}): {', '.join(dst.sheetnames)}")


def main() -> int:
    run()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
